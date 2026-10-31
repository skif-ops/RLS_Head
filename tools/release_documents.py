from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import re
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BASE_NAME = "EVT_ГСН_полный_производственный_комплект_2026-10-30.zip"
BASE_SHA = "a35f17e3f608441f8fe8eeaa38acc37e0d36a7cc6f1bc6d9fc4f3a4c6853af70"
VERSION = "1.0.1"
DATE = "2026-10-31"
NAME = f"EVT_ГСН_документы_v{VERSION}_{DATE}"
OUT = ROOT / "release" / NAME
ARCHIVE = ROOT / "release" / f"{NAME}.zip"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{" + NS["w"] + "}"
FONT = "Times New Roman"
INDENT_TWIPS = 709


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2) + "\n"


def baseline() -> zipfile.ZipFile:
    data = (ROOT / "release" / BASE_NAME).read_bytes()
    if digest(data) != BASE_SHA:
        raise ValueError("Baseline archive hash does not match the registered release")
    return zipfile.ZipFile(io.BytesIO(data))


def heading_style_ids(styles: ET.Element) -> set[str]:
    result = set()
    for style in styles.findall("w:style", NS):
        name = style.find("w:name", NS)
        label = name.get(W + "val", "") if name is not None else ""
        if re.fullmatch(r"heading ?[1-9]|заголовок ?[1-9]", label, re.IGNORECASE):
            result.add(style.get(W + "styleId", ""))
    return result


def heading_paragraphs(body: ET.Element, ids: set[str]) -> list[ET.Element]:
    return [p for p in body.iter(W + "p")
            if (s := p.find("w:pPr/w:pStyle", NS)) is not None
            and s.get(W + "val") in ids]


def child(parent: ET.Element, tag: str) -> ET.Element:
    element = parent.find(W + tag)
    if element is None:
        element = ET.Element(W + tag)
        if (parent.tag, tag) in {(W + "r", "rPr"), (W + "p", "pPr")}:
            parent.insert(0, element)
        elif parent.tag == W + "pPr" and tag == "ind":
            successors = {W + name for name in ("contextualSpacing", "mirrorIndents", "suppressOverlap",
                          "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl",
                          "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")}
            index = next((i for i, node in enumerate(parent) if node.tag in successors), len(parent))
            parent.insert(index, element)
        else:
            parent.append(element)
    return element


def font_properties(parent: ET.Element) -> None:
    fonts = child(child(parent, "rPr"), "rFonts")
    for key in list(fonts.attrib):
        if key.lower().endswith("theme"):
            del fonts.attrib[key]
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(W + key, FONT)


def indent_properties(parent: ET.Element) -> None:
    ind = child(child(parent, "pPr"), "ind")
    ind.attrib.clear()
    ind.set(W + "left", "0")
    ind.set(W + "right", "0")
    ind.set(W + "firstLine", str(INDENT_TWIPS))


def audit_docx(data: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        styles = ET.fromstring(archive.read("word/styles.xml"))
        body = ET.fromstring(archive.read("word/document.xml"))
    ids = heading_style_ids(styles)
    style_rows = []
    for style in styles.findall("w:style", NS):
        if style.get(W + "styleId") not in ids:
            continue
        fonts = style.find("w:rPr/w:rFonts", NS)
        ind = style.find("w:pPr/w:ind", NS)
        style_rows.append({
            "id": style.get(W + "styleId"),
            "fonts": {key: fonts.get(W + key) if fonts is not None else None
                      for key in ("ascii", "hAnsi", "eastAsia", "cs")},
            "font_theme": any(key.lower().endswith("theme") for key in fonts.attrib) if fonts is not None else False,
            "first_line_twips": ind.get(W + "firstLine") if ind is not None else None,
        })
    bad_paragraphs = 0
    paragraphs = heading_paragraphs(body, ids)
    for paragraph in paragraphs:
        ind = paragraph.find("w:pPr/w:ind", NS)
        if ind is None or ind.get(W + "firstLine") != str(INDENT_TWIPS):
            bad_paragraphs += 1
        for run in paragraph.iter(W + "r"):
            fonts = run.find("w:rPr/w:rFonts", NS)
            if fonts is None or any(fonts.get(W + key) != FONT for key in ("ascii", "hAnsi", "eastAsia", "cs")):
                bad_paragraphs += 1
    ok = all(row["first_line_twips"] == str(INDENT_TWIPS)
             and not row["font_theme"]
             and all(value == FONT for value in row["fonts"].values()) for row in style_rows)
    return {"status": "PASS" if ok and bad_paragraphs == 0 else "FAIL",
            "headings": len(paragraphs), "style_rows": style_rows,
            "paragraph_overrides_out_of_contract": bad_paragraphs}


def transform_docx(data: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    styles = ET.fromstring(source.read("word/styles.xml"))
    body = ET.fromstring(source.read("word/document.xml"))
    ids = heading_style_ids(styles)
    original_text = [item.text for item in body.iter(W + "t")]
    for style in styles.findall("w:style", NS):
        if style.get(W + "styleId") in {"Title", "Subtitle"}:
            font_properties(style)
        if style.get(W + "styleId") in ids:
            font_properties(style)
            indent_properties(style)
    for paragraph in heading_paragraphs(body, ids):
        indent_properties(paragraph)
        for run in paragraph.iter(W + "r"):
            font_properties(run)
    assert original_text == [item.text for item in body.iter(W + "t")]
    changed = {"word/styles.xml": ET.tostring(styles, encoding="utf-8", xml_declaration=True),
               "word/document.xml": ET.tostring(body, encoding="utf-8", xml_declaration=True)}
    core = etree.fromstring(source.read("docProps/core.xml"))
    core_ns = "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"
    version = core.find(core_ns + "version")
    if version is None:
        version = etree.SubElement(core, core_ns + "version")
    version.text = VERSION
    changed["docProps/core.xml"] = etree.tostring(core, encoding="UTF-8", xml_declaration=True, standalone=True)
    # Preserve namespace prefixes referenced by mc:Ignorable and field markup.
    for name in ("word/styles.xml", "word/document.xml"):
        original = etree.fromstring(source.read(name))
        updated = etree.fromstring(changed[name])
        for node in original:
            original.remove(node)
        original.extend(list(updated))
        changed[name] = etree.tostring(original, encoding="UTF-8", xml_declaration=True, standalone=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(copy.copy(info), changed.get(info.filename, source.read(info.filename)))
    result = output.getvalue()
    assert audit_docx(result)["status"] == "PASS"
    with zipfile.ZipFile(io.BytesIO(result)) as target:
        final_body = ET.fromstring(target.read("word/document.xml"))
        assert original_text == [item.text for item in final_body.iter(W + "t")]
        assert all(source.read(name) == target.read(name) for name in source.namelist() if name not in changed)
    return result


def new_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == "Normal" else 14 if name == "Title" else 12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.first_line_indent = Cm(1.25 if name != "Title" else 0)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.core_properties.version = VERSION
    doc.core_properties.subject = "Правила оформления и учета редакций документов"
    return doc


def administrative_documents() -> list[Path]:
    folder = OUT / "00_Версии_и_шаблон"
    folder.mkdir(parents=True)
    doc = new_doc()
    doc.add_paragraph("Учет версий документов", "Title")
    doc.add_paragraph(f"Выпуск документов v{VERSION} от 31.10.2026")
    sections = [
        ("1 Область выпуска", "Версия 1.0.1 содержит изменения оформления ранее выпущенных документов. Тексты технических разделов сохранены. Идентификаторы и инженерные редакции документов R1 сохраняются; версия комплекта ведется отдельно."),
        ("2 Обозначение версии", "Используется формат MAJOR.MINOR.PATCH. MAJOR повышается при несовместимой смене структуры или обозначений. MINOR повышается при добавлении или содержательном изменении документов. PATCH повышается при исправлении оформления и опечаток без изменения требований. Любое изменение выпущенного файла получает новую версию архива."),
        ("3 Имена и сохранность", "Имя архива включает состав, версию и дату: EVT_ГСН_документы_v1.0.1_2026-10-31.zip. Выпущенные архивы не перезаписываются. Исправленный выпуск получает новое имя. Дата и версия комплекта не заменяют редакцию платы, версию прошивки или номер экземпляра."),
        ("4 Исходный выпуск", "Архивы от 30.10.2026 с обозначением R1 зарегистрированы как исходный комплект 1.0.0. Это учетное обозначение введено задним числом; исходные имена, содержимое и контрольные суммы архивов сохранены."),
        ("5 Связь файлов", "RELEASE_REGISTRY.json связывает версии с датой, областью изменений, исходным коммитом, именами архивов и SHA-256. MANIFEST.json перечисляет содержимое нового архива и контрольные суммы каждого файла. Перед использованием сравнивают версию архива и запись реестра."),
        ("6 Оформление заголовков", "Заголовки разделов и подразделов используют Times New Roman и отступ первой строки 1,25 см от левого поля текста. В Word значение хранится как 709 twips, что соответствует примерно 1,2506 см. Левый и правый отступы абзаца равны нулю. Оформление титульных надписей по центру сохраняется."),
        ("7 Проверка выпуска", "Для каждого измененного DOCX проверяются свойства стилей и абзацев, сохранность исходного текста и визуальный результат. Архив проверяется на целостность. Контрольные суммы фиксируются после завершения оформления."),
    ]
    for title, text in sections:
        doc.add_heading(title, 1)
        doc.add_paragraph(text)
    policy = folder / f"Правила_учета_версий_документов_v{VERSION}.docx"
    doc.save(policy)
    policy.write_bytes(transform_docx(policy.read_bytes()))
    template = new_doc()
    template.add_paragraph("Наименование документа", "Title")
    template.add_paragraph("Обозначение документа: [заполнить]. Редакция: [заполнить].")
    template.add_paragraph("Версия комплекта: [заполнить]. Дата: [заполнить].")
    template.add_heading("1 Назначение документа", 1)
    template.add_paragraph("[Указать назначение и область применения документа.]")
    template.add_heading("1.1 Содержание раздела", 2)
    template.add_paragraph("[Внести текст подраздела.]")
    template.add_heading("1.1.1 Уточнение", 3)
    template.add_paragraph("[Внести текст пункта.]")
    template_path = folder / f"Шаблон_текстового_документа_v{VERSION}.docx"
    template.save(template_path)
    template_path.write_bytes(transform_docx(template_path.read_bytes()))
    return [policy, template_path]


def build() -> None:
    if OUT.exists() or ARCHIVE.exists():
        raise FileExistsError("Published output already exists; assign a new version")
    registry_path = ROOT / "release" / "RELEASE_REGISTRY.json"
    previous_registry = json.loads(registry_path.read_text(encoding="utf-8")) if registry_path.exists() else None
    if previous_registry and any(entry["version"] == VERSION for entry in previous_registry["entries"]):
        raise FileExistsError("Version already registered; assign a new version")
    if previous_registry:
        latest = tuple(map(int, previous_registry["latest_document_version"].split(".")))
        if tuple(map(int, VERSION.split("."))) <= latest:
            raise ValueError("New version must be greater than the registered version")
    OUT.mkdir(parents=True)
    report = []
    with baseline() as source:
        for name in sorted(source.namelist()):
            if not name.endswith(".docx"):
                continue
            source_relative = Path(*Path(name).parts[1:])
            relative = source_relative.with_name(f"{source_relative.stem}_v{VERSION}.docx")
            data = source.read(name)
            target = OUT / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            transformed = transform_docx(data)
            target.write_bytes(transformed)
            report.append({"path": relative.as_posix(), "source_path": source_relative.as_posix(), "source_sha256": digest(data),
                           "sha256": digest(transformed), "text_preserved": True,
                           "before": audit_docx(data), "after": audit_docx(transformed)})
    for path in administrative_documents():
        report.append({"path": path.relative_to(OUT).as_posix(), "new_administrative_document": True,
                       "after": audit_docx(path.read_bytes())})
    (OUT / "VERSION.txt").write_text(f"{VERSION}\n{DATE}\nОбновление оформления документов\n", encoding="utf-8")
    status = {"version": VERSION, "date": DATE, "scope": "DOCUMENT_FORMATTING_ONLY",
              "baseline_archive": BASE_NAME, "baseline_sha256": BASE_SHA,
              "engineering_document_revision": "R1", "new_bench_instruction_included": False,
              "production_authority": "NOT_AUTHORIZED", "documents": report}
    (OUT / "Проверка_оформления.json").write_text(json_text(status), encoding="utf-8")
    manifest = [{"path": path.relative_to(OUT).as_posix(), "bytes": path.stat().st_size,
                 "sha256": digest(path.read_bytes())}
                for path in sorted(OUT.rglob("*")) if path.is_file()]
    (OUT / "MANIFEST.json").write_text(json_text(manifest), encoding="utf-8")
    (OUT / "SHA256SUMS.txt").write_text("".join(f'{row["sha256"]}  {row["path"]}\n' for row in manifest), encoding="utf-8")
    zip_path = ARCHIVE
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as package:
        for path in sorted(OUT.rglob("*")):
            if path.is_file():
                package.write(path, f"{NAME}/{path.relative_to(OUT).as_posix()}")
    with zipfile.ZipFile(zip_path) as package:
        assert package.testzip() is None
    baseline_summary = json.loads((ROOT / "release" / "EVT_ГСН_ВЫПУСК_2026-10-30.json").read_text(encoding="utf-8"))
    current_commit = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    registry = {"schema": "DOCUMENT-RELEASE-REGISTRY-001", "latest_document_version": VERSION,
                "entries": [{"version": "1.0.0", "date": "2026-10-30", "scope": "LEGACY_BASELINE",
                             "retrospective_version_alias": True, "original_revision": "R1",
                             "source_commit": baseline_summary["source_commit"],
                             "archives": baseline_summary["archives"]},
                            {"version": VERSION, "date": DATE, "scope": "DOCUMENT_FORMATTING_ONLY",
                             "baseline_version": "1.0.0", "source_commit": current_commit,
                             "production_authority": "NOT_AUTHORIZED",
                             "archives": [{"name": zip_path.name, "bytes": zip_path.stat().st_size,
                                           "sha256": digest(zip_path.read_bytes())}]}]}
    if previous_registry:
        registry["entries"] = previous_registry["entries"] + registry["entries"][1:]
    registry_path.write_text(json_text(registry), encoding="utf-8")
    print(json_text({"documents": len(report), "updated_headings": sum(row["after"]["headings"] for row in report),
                     "archive": str(zip_path), "sha256": digest(zip_path.read_bytes())}))


def check() -> None:
    registry = json.loads((ROOT / "release" / "RELEASE_REGISTRY.json").read_text(encoding="utf-8"))
    for entry in registry["entries"]:
        for item in entry["archives"]:
            data = (ROOT / "release" / item["name"]).read_bytes()
            assert digest(data) == item["sha256"], item["name"]
            assert len(data) == item["bytes"], item["name"]
    manifest = json.loads((OUT / "MANIFEST.json").read_text(encoding="utf-8"))
    for item in manifest:
        data = (OUT / item["path"]).read_bytes()
        assert digest(data) == item["sha256"], item["path"]
        assert len(data) == item["bytes"], item["path"]
    with zipfile.ZipFile(ARCHIVE) as package:
        assert package.testzip() is None
        for path in OUT.rglob("*"):
            if path.is_file():
                name = f"{NAME}/{path.relative_to(OUT).as_posix()}"
                assert package.read(name) == path.read_bytes(), name
    documents = list(OUT.rglob("*.docx"))
    assert len(documents) == 13
    for path in documents:
        assert audit_docx(path.read_bytes())["status"] == "PASS", path.name
        assert Document(path).core_properties.version == VERSION, path.name
    print(json_text({"status": "PASS", "documents": len(documents),
                     "manifest_entries": len(manifest), "registered_archives": sum(len(e["archives"]) for e in registry["entries"])}))


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit or release document formatting without changing technical content")
    parser.add_argument("operation", choices=["audit", "build", "check"])
    args = parser.parse_args()
    if args.operation == "build":
        build()
    elif args.operation == "check":
        check()
    else:
        with baseline() as source:
            print(json_text([{ "path": name, **audit_docx(source.read(name))}
                             for name in sorted(source.namelist()) if name.endswith(".docx")]))


if __name__ == "__main__":
    main()
