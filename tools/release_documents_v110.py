from __future__ import annotations

import argparse
import hashlib
import io
import json
import shutil
import subprocess
import zipfile
from pathlib import Path, PurePosixPath

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from tools.docx_style_policy import audit_docx, transform_docx


ROOT = Path(__file__).resolve().parents[1]
RELEASE = ROOT / "release"
SOURCE_ARCHIVE = RELEASE / "EVT_ГСН_документы_v1.0.1_2026-10-31.zip"
SOURCE_SHA256 = "165dde82aa381e87357110661f28acef0c944520ac41babe26913ee36a0397c6"
VERSION = "1.1.0"
DATE = "2026-10-31"
PACKAGE_NAME = f"EVT_ГСН_документы_v{VERSION}_{DATE}"
PACKAGE_DIR = RELEASE / PACKAGE_NAME
PACKAGE_ARCHIVE = RELEASE / f"{PACKAGE_NAME}.zip"
METHODS_NAME = f"EVT_ГСН_ПМ_EVT_и_отчетные_материалы_v{VERSION}_{DATE}"
METHODS_ARCHIVE = RELEASE / f"{METHODS_NAME}.zip"
REGISTRY_PATH = RELEASE / "RELEASE_REGISTRY.json"
TEMPLATE_TARGET = ROOT / "templates" / f"Текстовый_документ_v{VERSION}.docx"


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2) + "\n"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_table_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table_width = sum(widths_cm)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(Cm(table_width).twips)))
    tbl_w.set(qn("w:type"), "dxa")
    for index, width in enumerate(widths_cm):
        table.columns[index].width = Cm(width)
        grid_columns = table._tbl.tblGrid.gridCol_lst
        if index < len(grid_columns):
            grid_columns[index].set(qn("w:w"), str(int(Cm(width).twips)))
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)
        for index, width in enumerate(widths_cm):
            cell = row.cells[index]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(Cm(width).twips)))
            tc_w.set(qn("w:type"), "dxa")


def fill_table(
    table,
    headers: list[str],
    rows: list[list[str]],
    widths_cm: list[float] | None = None,
) -> None:
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        header.cells[index].text = text
        set_cell_shading(header.cells[index], "D9EAF7")
        header.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in header.cells[index].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for values in rows:
        cells = table.add_row().cells
        for index, text in enumerate(values):
            cells[index].text = text
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if widths_cm is not None:
        set_table_widths(table, widths_cm)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def equipment_rows() -> list[list[str]]:
    return [
        ["Лабораторный источник питания", "0-15 В, не менее 5 А, установка ограничения тока, индикация тока не хуже 10 мА"],
        ["Цифровой мультиметр", "Измерение сопротивления и постоянного напряжения; щупы с изолированными наконечниками"],
        ["Осциллограф", "Не менее 100 МГц, щупы 10:1, пружинный контакт земли, ограничение полосы 20 МГц"],
        ["Тепловизор или термопара", "Диапазон не уже 20-100 °C, регистрация температуры во времени"],
        ["Программатор", "ST-LINK и кабель Tag-Connect для J10; питание платы от ST-LINK не подавать"],
        ["CAN-FD интерфейс", "500 кбит/с nominal, 2,5 Мбит/с data, BRS, стандартные идентификаторы, кадр 64 байта"],
        ["Два адаптера USB-UART", "Только уровни 3,3 В, 921600 бод, 8 бит, без четности, 1 стоп-бит"],
        ["Генератор импульсов", "Уровни 0-3,3 В, выход с общей землей; для RADAR_REF и GNSS_PPS"],
        ["Средства регистрации", "Компьютер с STM32CubeProgrammer, UART логгером, CAN-FD ПО и декодером TargetState V1"],
    ]


def connector_rows() -> list[list[str]]:
    return [
        ["J1", "1: 12V_IN; 2: GND", "Вход питания платы, 12,0 В"],
        ["J5", "1: RADAR_12V; 2: GND", "Защищенный выход питания внешнего радара"],
        ["J4", "1: CANL; 2: GND; 3: CANH; 4: RADAR_REF; 5: GND; 6: NC", "CAN-FD и опорный импульс радара"],
        ["J6", "1: GND; 2: UART1_TX; 3: UART1_RX; 4: GND; 5: UART2_TX; 6: UART2_RX", "Два независимых UART 3,3 В"],
        ["J8", "1: 5V_GNSS; 2: GNSS_RX_EXT; 3: GNSS_TX_EXT; 4: GNSS_PPS; 5: NC; 6: GND", "Съемный внешний GNSS модуль"],
        ["J10", "1: VREF 3V3; 2: SWDIO; 3: GND; 4: SWCLK; 5: GND; 6: SWO", "Программирование и отладка"],
    ]


def step_rows() -> list[list[str]]:
    return [
        ["1", "Идентифицировать образец", "Записаны серийный номер платы, ревизия Pilot R1, версия комплекта, имя и SHA-256 прошивки.", "Если маркировка платы или файл прошивки не соответствуют описи."],
        ["2", "Выполнить внешний осмотр без питания", "Нет трещин, загрязнений, непропая, перемычек, нарушенной полярности. Разъемы и крепеж не повреждены. R40 и R70 не установлены по базовой комплектации.", "При любом повреждении, перемычке припоя, неверной полярности или неоговоренном отклонении монтажа."],
        ["3", "Проверить сопротивления", "После выдержки щупов 30 с нет устойчивого сопротивления менее 2 Ом между GND и J1.1, J5.1, J8.1, J10.1. Значения записаны как исходная база экземпляра.", "Если значение стабилизируется ниже 2 Ом или уменьшается со временем. Питание не подавать."],
        ["4", "Подготовить первое включение", "Внешние радар, GNSS, UART, CAN и ST-LINK отключены. Источник выключен, установлен 12,0 В и предел 0,50 А. Плюс идет на J1.1, минус на J1.2.", "Если полярность или назначение контактов нельзя однозначно подтвердить."],
        ["5", "Подать питание на 2 с", "Плата не достигает ограничения 0,50 А, нет дыма, запаха, искрения, акустического шума и резкого нагрева.", "Немедленно выключить при любом признаке повреждения или работе ограничения тока."],
        ["6", "Измерить шины в первые 10 с", "J10.1: 3,135-3,465 В; J8.1: 4,75-5,25 В; J5.1: 11,4-12,6 В относительно ближайшего GND.", "Немедленно выключить при обратной полярности, выходе любой шины за пределы или периодическом провале напряжения."],
        ["7", "Измерить пульсации", "При полосе 20 МГц и короткой земле: 3V3 и 5V_GNSS не более 100 мВ пик-пик; RADAR_12V не более 200 мВ пик-пик без внешней нагрузки.", "Выключить при колебаниях преобразователя, выбросах более 200 мВ на 3V3/5V_GNSS или более 500 мВ на RADAR_12V."],
        ["8", "Выдержать плату 15 мин без внешних модулей", "Потребление не достигает 0,50 А. Температура любого корпуса не выше 70 °C и после 5-й минуты не растет быстрее 5 °C/мин.", "Выключить при 80 °C, росте температуры без стабилизации, сбросах питания или срабатывании eFuse."],
        ["9", "Подключить ST-LINK при снятом питании", "J10 подключен по распиновке. После подачи 12 В программатор видит VREF 3,135-3,465 В и распознает STM32H753.", "Не продолжать при неверном VREF, невозможности стабильного подключения или попытке питания платы от программатора."],
        ["10", "Записать и проверить прошивку", "Записан утвержденный HEX, verify завершен без ошибок, лог и SHA-256 сохранены, после reset процессор работает без повторяющегося сброса.", "При ошибке erase, program, verify, неожиданной защите памяти или циклическом reset."],
        ["11", "Проверить структуру g_evt_app_status", "magic = 0x53545645; version = 1.1; init_error_flags = 0; config_crc32c не равен 0; boot_id не равен 0; imu_who_am_i = 0x47; imu_samples растет; imu_errors не растет.", "При любом init_error_flags, неверном WHO_AM_I, нулевом CRC конфигурации или росте imu_errors."],
        ["12", "Проверить состояние без внешних источников", "Допустимы признаки RADAR_STALE и GNSS_PPS_STALE. Признак POWER_FAULT отсутствует. После запуска IMU признак IMU_STALE отсутствует.", "При POWER_FAULT, устойчивом IMU_STALE или изменении fault flags без воздействия."],
        ["13", "Проверить оба UART без подачи сигналов на входы", "На J6.2 и J6.5 уровень покоя 3,0-3,3 В. На входы J6.3 и J6.6 не поступает более 3,3 В. До приема полного HIL пакета передача TargetState не требуется.", "При уровне выше 3,6 В, устойчивом низком уровне TX или заметной нагрузке через адаптер."],
        ["14", "Подключить CAN-FD при снятом питании", "J4.1 к CANL, J4.3 к CANH, J4.2 к общей земле. На стенде есть корректная оконечная нагрузка. Настройки: 500 кбит/с, 2,5 Мбит/с, BRS, ID 0x320, 64 байта.", "При перепутанных CANH/CANL, отсутствии общей земли или неизвестной оконечной нагрузке."],
        ["15", "Передать контрольную HIL-R1 запись", "radar_packets и targetstate_packets растут. transport_rejects, can_bus_offs, uart1_drops и uart2_drops не растут. Оба UART выдают пакеты по 200 байт.", "При bus-off, росте rejects/drops, отсутствии выдачи на одном канале или различии смысловых данных двух UART."],
        ["16", "Декодировать TargetState V1", "Magic 0x31474C52, версия 1.0, тип 1, длина 200 байт, CRC32C PASS, sequence монотонна, boot_id совпадает с g_evt_app_status, значения валидных float конечны.", "При неверной длине, magic, версии, CRC, повторе sequence, NaN или бесконечности в валидном поле."],
        ["17", "Проверить RADAR_REF", "Подать на J4.4 импульсы 0-3,3 В частотой 10 Гц относительно J4.5. Не позднее 250 мс после первого импульса RADAR_STALE снимается и не возникает при непрерывной серии.", "При амплитуде выше 3,6 В, пропусках, повторном RADAR_STALE или нарушении работы CAN/UART."],
        ["18", "Проверить внешний GNSS", "Подключать J8 только при снятом питании. Модуль получает 4,75-5,25 В, USART3 работает 115200 8N1, PPS имеет уровни 0-3,3 В. При валидных RMC/GGA и PPS снимаются признаки UART_STALE, PPS_STALE, FIX_INVALID и TIME_INVALID.", "При токе модуля выше 0,50 А, просадке 5V_GNSS, уровне сигнала выше 3,6 В, переполюсовке или росте gnss_overflows."],
        ["19", "Проверить потерю GNSS", "После отключения данных или PPS обработка радара и оба UART продолжаются. Не позднее 1,5 с появляется stale/holdover, TIME_VALID снимается или неопределенность времени растет.", "При остановке выдачи TargetState, зависании, сбросе или ошибочном сохранении TIME_VALID без роста неопределенности."],
        ["20", "Проверить watchdog", "Контролируемо прекратить refresh более чем на 2 с в отдельном тестовом сеансе. Происходит reset, boot_id меняется, после запуска интерфейсы восстанавливаются.", "При отсутствии reset, циклическом reset или невозможности восстановить штатный запуск."],
        ["21", "Завершить проверку", "Сохранены фото, измерения, логи SWD, CAN и UART, температуры, идентификаторы оборудования, версии ПО и статусы всех шагов.", "Комплект не переводить в PASS при отсутствии хотя бы одного обязательного подтверждения."],
    ]


def add_bench_procedure(doc: Document, include_title: bool = False) -> None:
    if include_title:
        doc.add_paragraph("ПЕРВОЕ ВКЛЮЧЕНИЕ И СТЕНДОВАЯ ПРОВЕРКА ПИЛОТНОГО ОБРАЗЦА", style="Title")
        add_paragraph(doc, "Обозначение: EVT-GSN-PWRON-001. Редакция: R1. Версия комплекта: 1.1.0.")
    add_heading(doc, "1 Назначение и правило остановки" if include_title else "7 Пошаговая инструкция первого включения и стендовой проверки", 1)
    add_paragraph(doc, "Методика применяется к плате EVT_SYS_CARRIER Pilot R1 до подключения к остальным узлам стенда. Каждый шаг выполняют только после положительного результата предыдущего шага. При выполнении любого критерия остановки питание немедленно снимают, состояние стенда фотографируют, записывают показания и открывают запись об отклонении. Повторное включение допускается после установления причины и письменного разрешения ответственного за EVT.")
    add_paragraph(doc, "Указанные пределы предназначены для первого пилотного образца и стендовой проверки. Они не выдают разрешение на серийное производство. Production authority сохраняется в состоянии NOT_AUTHORIZED.")

    add_heading(doc, "2 Оборудование" if include_title else "7.1 Оборудование", 2)
    table = doc.add_table(rows=1, cols=2)
    fill_table(table, ["Средство", "Требование"], equipment_rows(), [7.0, 18.0])

    add_heading(doc, "3 Контрольные разъемы" if include_title else "7.2 Контрольные разъемы", 2)
    table = doc.add_table(rows=1, cols=3)
    fill_table(table, ["Разъем", "Контакты", "Назначение"], connector_rows(), [3.0, 13.0, 9.0])
    add_paragraph(doc, "Для всех измерений напряжения используют ближайший указанный контакт GND. Входы UART, RADAR_REF и GNSS_PPS допускают только уровни 3,3 В. Подавать 5 В на сигнальные контакты запрещено.")

    add_heading(doc, "4 Последовательность проверки" if include_title else "7.3 Последовательность проверки", 2)
    table = doc.add_table(rows=1, cols=4)
    fill_table(
        table,
        ["Шаг", "Действие", "Ожидаемый результат", "Критерий остановки"],
        step_rows(),
        [1.2, 4.8, 9.6, 9.4],
    )

    add_heading(doc, "5 Запись результатов" if include_title else "7.4 Запись результатов", 2)
    add_paragraph(doc, "Для каждого шага фиксируют PASS, FAIL, BLOCKED или NOT RUN, фактические значения, серийный номер образца, дату, исполнителя и путь к evidence. Значение не заменяют словами «норма» или «соответствует». Для напряжения, тока, пульсаций и температуры записывают число, единицу измерения и точку измерения.")
    add_paragraph(doc, "Итог FIRST POWER-ON PASS допускается только при PASS шагов 1-13. Итог BENCH FUNCTION PASS допускается только при PASS шагов 1-21. BLOCKED или NOT RUN не приравниваются к PASS. Любой незакрытый FAIL сохраняет общий результат FAIL.")

    add_heading(doc, "6 Минимальный протокол" if include_title else "7.5 Минимальный протокол", 2)
    rows = [
        ["Идентификация", "Серийный номер платы, ревизия, версия комплекта, commit, SHA-256 прошивки"],
        ["Питание", "Сопротивления до включения, ток 2 с/1 мин/15 мин, 3V3, 5V_GNSS, RADAR_12V, пульсации"],
        ["Температура", "Термограммы или значения 1 мин, 5 мин и 15 мин с указанием горячих компонентов"],
        ["Прошивка", "Лог program/verify, значения g_evt_app_status до и после интерфейсных тестов"],
        ["Интерфейсы", "CAN-FD настройки и счетчики, UART1/UART2 файлы, результаты TargetState decode и CRC32C"],
        ["Время и GNSS", "Частота/уровень RADAR_REF, источник GNSS/PPS, переход в holdover, fault flags"],
        ["Решение", "PASS/FAIL/BLOCKED/NOT RUN по шагам, перечень отклонений, подписи исполнителя и проверяющего"],
    ]
    table = doc.add_table(rows=1, cols=2)
    fill_table(table, ["Раздел", "Обязательная запись"], rows, [6.0, 19.0])


def make_section_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def make_section_a3_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(42)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def make_document_landscape(data: bytes) -> bytes:
    doc = Document(io.BytesIO(data))
    for section in doc.sections:
        make_section_landscape(section)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def prepare_wide_source_document(data: bytes, name: str) -> bytes:
    doc = Document(io.BytesIO(data))
    for section in doc.sections:
        if name.startswith((
            "Схема_соединений_и_задание_на_жгуты",
            "Форма_отчета_EVT",
        )):
            make_section_a3_landscape(section)
        else:
            make_section_landscape(section)
    for table in doc.tables:
        columns = len(table.columns)
        if name.startswith("Протокол_TargetState_V1"):
            if columns == 5:
                set_table_widths(table, [2.3, 2.3, 8.2, 3.4, 8.8])
            elif columns == 3:
                set_table_widths(table, [3.5, 11.0, 10.5])
            elif columns == 4:
                set_table_widths(table, [4.0, 5.0, 8.0, 8.0])
        elif name.startswith("Схема_соединений_и_задание_на_жгуты"):
            if columns == 5:
                set_table_widths(table, [4.0, 3.0, 8.0, 11.0, 11.5])
            elif columns == 6:
                set_table_widths(table, [3.0, 2.0, 6.5, 8.0, 7.5, 10.5])
            elif columns == 4:
                set_table_widths(table, [4.5, 6.0, 12.0, 15.0])
        elif name.startswith("Форма_отчета_EVT"):
            if columns == 5:
                set_table_widths(table, [4.5, 8.5, 4.0, 7.5, 13.0])
            elif columns == 2:
                set_table_widths(table, [12.0, 25.5])
            elif columns == 4:
                set_table_widths(table, [4.5, 10.5, 7.5, 15.0])
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def new_document(landscape: bool = False) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    if landscape:
        make_section_landscape(section)
    for style in doc.styles:
        if style.type in {1, 2}:
            style.font.name = "Times New Roman"
    doc.core_properties.version = VERSION
    return doc


def save_document(doc: Document) -> bytes:
    stream = io.BytesIO()
    doc.save(stream)
    result, _ = transform_docx(stream.getvalue(), version=VERSION)
    return result


def policy_document() -> bytes:
    doc = new_document()
    doc.add_paragraph("ТРЕБОВАНИЯ К ОФОРМЛЕНИЮ И УЧЕТУ ВЕРСИЙ ДОКУМЕНТОВ", style="Title")
    add_paragraph(doc, f"Версия требований: {VERSION}. Дата: {DATE}.")
    sections = [
        ("1 Область применения", "Требования применяют ко всем последующим документам DOCX во всех проектах. Если нормативный документ или прямо утвержденный заказчиком шаблон устанавливает более строгое правило, применяют более строгое правило и фиксируют исключение."),
        ("2 Основной текст", "Шрифт Times New Roman, размер 14 пт, одинарный межстрочный интервал. Абзац выравнивают по ширине. Отступ первой строки составляет 1,25 см. Интервалы до и после абзаца равны 0 пт."),
        ("3 Таблицы", "Текст в таблицах оформляют шрифтом Times New Roman, 12 пт, с одинарным межстрочным интервалом и выравниванием по ширине. Отступ первой строки в ячейках не применяют. Заголовочную строку таблицы выделяют полужирным шрифтом и повторяют на следующих страницах, если таблица переносится."),
        ("4 Заголовки", "Заголовки оформляют шрифтом Times New Roman, 14 пт, полужирным начертанием, с одинарным межстрочным интервалом, выравниванием по ширине и отступом первой строки 1,25 см. Интервалы до и после заголовка равны 0 пт."),
        ("5 Знаки препинания", "Длинное тире и среднее тире не употребляют. Для смыслового разделения используют дефис, двоеточие, точку с запятой или перестраивают предложение."),
        ("6 Проверка DOCX", "Перед выпуском автоматически проверяют шрифт и размер каждого текстового фрагмента, абзацные отступы, межстрочный интервал, выравнивание, полужирное начертание заголовков, отсутствие длинного и среднего тире. Затем документ преобразуют в PDF или постраничные изображения и проверяют визуально."),
        ("7 Версионность", "Версия имеет формат MAJOR.MINOR.PATCH. MAJOR изменяют при несовместимой смене структуры или обозначений. MINOR изменяют при добавлении методики или другом содержательном изменении. PATCH изменяют при исправлении оформления и опечаток без изменения требований. Выпущенный архив не перезаписывают."),
        ("8 Идентификация выпуска", "Версию указывают в имени архива, имени документа и свойствах DOCX. Реестр выпуска содержит дату, исходный commit, область изменений, SHA-256 и связь с предыдущей версией. Версия комплекта не заменяет ревизию платы или версию прошивки."),
    ]
    for title, text in sections:
        add_heading(doc, title, 1)
        add_paragraph(doc, text)
    return save_document(doc)


def template_document() -> bytes:
    doc = new_document()
    doc.add_paragraph("НАИМЕНОВАНИЕ ДОКУМЕНТА", style="Title")
    add_paragraph(doc, "Обозначение: [заполнить]. Редакция: [заполнить]. Версия комплекта: [заполнить].")
    add_heading(doc, "1 Назначение", 1)
    add_paragraph(doc, "[Указать назначение и область применения документа.]")
    add_heading(doc, "1.1 Исходные данные", 2)
    add_paragraph(doc, "[Перечислить утвержденные исходные данные и версии источников.]")
    add_heading(doc, "1.1.1 Порядок применения", 3)
    add_paragraph(doc, "[Привести конкретный порядок действий и критерии результата.]")
    table = doc.add_table(rows=1, cols=3)
    fill_table(table, ["Параметр", "Требование", "Результат"], [["[заполнить]", "[заполнить]", "[заполнить]"]])
    return save_document(doc)


def standalone_instruction() -> bytes:
    doc = new_document(landscape=True)
    add_bench_procedure(doc, include_title=True)
    return save_document(doc)


def append_instruction(data: bytes) -> bytes:
    doc = Document(io.BytesIO(data))
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    make_section_landscape(section)
    add_bench_procedure(doc, include_title=False)
    return save_document(doc)


def source_documents() -> dict[str, bytes]:
    payload = SOURCE_ARCHIVE.read_bytes()
    if digest(payload) != SOURCE_SHA256:
        raise ValueError("Source document archive does not match the registered SHA-256")
    result = {}
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        for name in archive.namelist():
            if not name.endswith(".docx"):
                continue
            relative = PurePosixPath(name)
            if len(relative.parts) < 2:
                continue
            result[PurePosixPath(*relative.parts[1:]).as_posix()] = archive.read(name)
    if len(result) != 13:
        raise ValueError(f"Expected 13 source DOCX files, got {len(result)}")
    return result


def output_name(relative: str) -> Path:
    path = Path(*PurePosixPath(relative).parts)
    stem = path.stem.removesuffix("_v1.0.1")
    return path.with_name(f"{stem}_v{VERSION}.docx")


def write_manifest(folder: Path) -> list[dict]:
    manifest = [
        {
            "path": path.relative_to(folder).as_posix(),
            "bytes": path.stat().st_size,
            "sha256": digest(path.read_bytes()),
        }
        for path in sorted(folder.rglob("*"))
        if path.is_file()
    ]
    (folder / "MANIFEST.json").write_text(json_text(manifest), encoding="utf-8")
    (folder / "SHA256SUMS.txt").write_text(
        "".join(f'{row["sha256"]}  {row["path"]}\n' for row in manifest),
        encoding="utf-8",
    )
    return manifest


def make_zip(folder: Path, archive_path: Path, root_name: str) -> None:
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(folder.rglob("*")):
            if path.is_file():
                archive.write(path, f"{root_name}/{path.relative_to(folder).as_posix()}")
    with zipfile.ZipFile(archive_path) as archive:
        if archive.testzip() is not None:
            raise ValueError(f"Damaged archive: {archive_path.name}")


def build_methods_archive() -> None:
    method_files = {
        "00_Программа_и_методики/Программа_и_методики_EVT_v1.1.0.docx": PACKAGE_DIR / "08_Контроль_и_приемка" / "Программа_и_методики_EVT_v1.1.0.docx",
        "00_Программа_и_методики/Инструкция_первого_включения_и_стендовой_проверки_v1.1.0.docx": PACKAGE_DIR / "08_Контроль_и_приемка" / "Инструкция_первого_включения_и_стендовой_проверки_v1.1.0.docx",
        "01_Формы_отчетов/Форма_отчета_EVT_v1.1.0.docx": PACKAGE_DIR / "10_Отчетные_материалы" / "Форма_отчета_EVT_v1.1.0.docx",
    }
    with zipfile.ZipFile(METHODS_ARCHIVE, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        hashes = []
        for relative, source in method_files.items():
            data = source.read_bytes()
            archive.writestr(f"{METHODS_NAME}/{relative}", data)
            hashes.append({"path": relative, "bytes": len(data), "sha256": digest(data)})
        archive.writestr(
            f"{METHODS_NAME}/VERSION.txt",
            f"{VERSION}\n{DATE}\nСодержательное дополнение ПМ и унификация оформления DOCX\n",
        )
        archive.writestr(
            f"{METHODS_NAME}/MANIFEST.json",
            json_text(hashes),
        )
    with zipfile.ZipFile(METHODS_ARCHIVE) as archive:
        if archive.testzip() is not None:
            raise ValueError("Damaged methods archive")


def update_registry() -> None:
    registry = json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))
    if any(entry["version"] == VERSION for entry in registry["entries"]):
        raise FileExistsError(f"Version {VERSION} is already registered")
    current_commit = subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
    ).strip()
    registry["latest_document_version"] = VERSION
    registry["entries"].append(
        {
            "version": VERSION,
            "date": DATE,
            "scope": "DOCUMENT_STYLE_AND_EVT_METHOD_UPDATE",
            "baseline_version": "1.0.1",
            "source_commit": current_commit,
            "changes": [
                "Полная унификация оформления всех DOCX",
                "Добавлена пошаговая методика первого включения и стендовой проверки",
                "Требования оформления приняты для последующих документов",
            ],
            "production_authority": "NOT_AUTHORIZED",
            "archives": [
                {
                    "name": PACKAGE_ARCHIVE.name,
                    "bytes": PACKAGE_ARCHIVE.stat().st_size,
                    "sha256": digest(PACKAGE_ARCHIVE.read_bytes()),
                },
                {
                    "name": METHODS_ARCHIVE.name,
                    "bytes": METHODS_ARCHIVE.stat().st_size,
                    "sha256": digest(METHODS_ARCHIVE.read_bytes()),
                },
            ],
        }
    )
    REGISTRY_PATH.write_text(json_text(registry), encoding="utf-8")


def build() -> None:
    if PACKAGE_DIR.exists() or PACKAGE_ARCHIVE.exists() or METHODS_ARCHIVE.exists():
        raise FileExistsError("Versioned output already exists")
    latest = json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))["latest_document_version"]
    if tuple(map(int, latest.split("."))) >= tuple(map(int, VERSION.split("."))):
        raise ValueError("Version must be greater than the latest registered version")

    PACKAGE_DIR.mkdir(parents=True)
    reports = []
    for relative, data in sorted(source_documents().items()):
        target_relative = output_name(relative)
        if "00_Версии_и_шаблон" in target_relative.parts:
            continue
        if target_relative.name.startswith("Программа_и_методики_EVT"):
            transformed = append_instruction(data)
            change = "APPENDED_FIRST_POWER_ON_AND_BENCH_PROCEDURE"
        else:
            if target_relative.name.startswith((
                "Протокол_TargetState_V1",
                "Схема_соединений_и_задание_на_жгуты",
                "Форма_отчета_EVT",
            )):
                data = prepare_wide_source_document(data, target_relative.name)
            transformed, _ = transform_docx(data, version=VERSION)
            change = "STYLE_NORMALIZATION"
        target = PACKAGE_DIR / target_relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(transformed)
        reports.append({
            "path": target_relative.as_posix(),
            "source_path": relative,
            "source_sha256": digest(data),
            "sha256": digest(transformed),
            "change": change,
            "audit": audit_docx(transformed),
        })

    policy = PACKAGE_DIR / "00_Версии_и_шаблон" / f"Правила_оформления_и_версий_v{VERSION}.docx"
    template = PACKAGE_DIR / "00_Версии_и_шаблон" / f"Шаблон_текстового_документа_v{VERSION}.docx"
    instruction = PACKAGE_DIR / "08_Контроль_и_приемка" / f"Инструкция_первого_включения_и_стендовой_проверки_v{VERSION}.docx"
    for path, payload, change in (
        (policy, policy_document(), "REPLACED_STYLE_POLICY"),
        (template, template_document(), "REPLACED_TEMPLATE"),
        (instruction, standalone_instruction(), "NEW_EVT_METHOD"),
    ):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload)
        reports.append({
            "path": path.relative_to(PACKAGE_DIR).as_posix(),
            "sha256": digest(payload),
            "change": change,
            "audit": audit_docx(payload),
        })

    TEMPLATE_TARGET.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, TEMPLATE_TARGET)
    (PACKAGE_DIR / "VERSION.txt").write_text(
        f"{VERSION}\n{DATE}\nСодержательное дополнение ПМ и унификация оформления DOCX\n",
        encoding="utf-8",
    )
    qa = {
        "version": VERSION,
        "date": DATE,
        "baseline_version": "1.0.1",
        "source_archive": SOURCE_ARCHIVE.name,
        "source_sha256": SOURCE_SHA256,
        "production_authority": "NOT_AUTHORIZED",
        "documents": reports,
    }
    (PACKAGE_DIR / "Проверка_оформления.json").write_text(
        json_text(qa), encoding="utf-8"
    )
    write_manifest(PACKAGE_DIR)
    make_zip(PACKAGE_DIR, PACKAGE_ARCHIVE, PACKAGE_NAME)
    build_methods_archive()
    update_registry()
    print(json_text({
        "status": "BUILT",
        "version": VERSION,
        "documents": len(reports),
        "document_archive": PACKAGE_ARCHIVE.name,
        "document_sha256": digest(PACKAGE_ARCHIVE.read_bytes()),
        "methods_archive": METHODS_ARCHIVE.name,
        "methods_sha256": digest(METHODS_ARCHIVE.read_bytes()),
    }))


def check() -> None:
    registry = json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))
    for entry in registry["entries"]:
        for archive in entry["archives"]:
            data = (RELEASE / archive["name"]).read_bytes()
            if digest(data) != archive["sha256"] or len(data) != archive["bytes"]:
                raise ValueError(f"Registry mismatch: {archive['name']}")
    documents = sorted(PACKAGE_DIR.rglob("*.docx"))
    if len(documents) != 14:
        raise ValueError(f"Expected 14 DOCX files, got {len(documents)}")
    failures = []
    for path in documents:
        audit = audit_docx(path.read_bytes())
        if audit["status"] != "PASS":
            failures.append({"path": str(path), "audit": audit})
    with zipfile.ZipFile(PACKAGE_ARCHIVE) as archive:
        if archive.testzip() is not None:
            failures.append({"archive": PACKAGE_ARCHIVE.name})
        for path in PACKAGE_DIR.rglob("*"):
            if path.is_file():
                name = f"{PACKAGE_NAME}/{path.relative_to(PACKAGE_DIR).as_posix()}"
                if archive.read(name) != path.read_bytes():
                    failures.append({"archive_path": name})
    with zipfile.ZipFile(METHODS_ARCHIVE) as archive:
        if archive.testzip() is not None:
            failures.append({"archive": METHODS_ARCHIVE.name})
    if failures:
        raise ValueError(json_text(failures))
    print(json_text({
        "status": "PASS",
        "documents": len(documents),
        "paragraphs": sum(audit_docx(path.read_bytes())["paragraphs"] for path in documents),
        "table_paragraphs": sum(audit_docx(path.read_bytes())["table_paragraphs"] for path in documents),
        "forbidden_dashes": sum(audit_docx(path.read_bytes())["forbidden_dashes"] for path in documents),
        "registered_archives": sum(len(entry["archives"]) for entry in registry["entries"]),
    }))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("operation", choices=("build", "check"))
    args = parser.parse_args()
    if args.operation == "build":
        build()
    else:
        check()


if __name__ == "__main__":
    main()
